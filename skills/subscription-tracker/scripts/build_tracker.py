#!/usr/bin/env python3
"""Render a Subscription Tracker snapshot from a list of subscription rows.

This is a thin, dependency-light helper for the `subscription-tracker` skill. It
takes the rows you assembled from email (as a JSON list of objects) and writes a
dated snapshot spreadsheet. It exists so the column schema and the value rules
live in one place instead of being re-typed by hand each run.

Design notes that match the skill:
- The schema is fixed and ordered. Do not reorder columns; downstream reviews
  and prior snapshots assume this order.
- `Auto-Renew` and `Status` are two INDEPENDENT dimensions. `Next Event` means
  different things depending on `Auto-Renew` (charge date vs access-end date vs
  blank). This script does not "correct" them for you — it renders what you give
  it — but it validates the enum values so a typo can't quietly poison a report.
- Unknown amounts must be the literal string `VERIFY`, never a guessed number,
  so they never get summed into spend totals.

Output format:
- Writes `.xlsx` when `openpyxl` is available (preferred — that's the local copy
  the skill delivers).
- Falls back to `.csv` when it isn't, so this always produces *something* even in
  a bare environment. The cloud snapshot can be created from either.

Usage:
    python build_tracker.py --rows rows.json --date 2026-07-15 [--outdir .]
    cat rows.json | python build_tracker.py --date 2026-07-15

`rows.json` is a JSON array of objects keyed by the column names below (any
missing key renders blank). Example row:

    {
      "Vendor": "Netflix",
      "Category": "Media",
      "Plan/Tier": "Standard",
      "Amount": "15.49",
      "Billing Cycle": "Monthly",
      "Last Charge": "2026-07-02",
      "Auto-Renew": "On",
      "Next Event": "2026-08-02",
      "Status": "Active",
      "Notes": "next date estimated from cycle",
      "Last Updated": "2026-07-15"
    }
"""

import argparse
import csv
import json
import sys

# Fixed, ordered schema. Keep in sync with SKILL.md "Columns".
COLUMNS = [
    "Vendor",
    "Category",
    "Plan/Tier",
    "Amount",
    "Billing Cycle",
    "Last Charge",
    "Auto-Renew",
    "Next Event",
    "Status",
    "Notes",
    "Last Updated",
]

# Small, generic buckets so the user can group spend (see SKILL.md).
CATEGORIES = {
    "Software", "Media", "Phone", "Health", "Home",
    "Professional", "Finance", "Other",
}

AUTO_RENEW_VALUES = {"On", "Off", "Verify"}
STATUS_VALUES = {"Active", "Ended", "Verify"}


def validate(rows):
    """Return a list of human-readable warnings; never raises on data issues.

    We warn rather than hard-fail so a single malformed row can't block the whole
    snapshot — but the warnings must surface, because in a money task a silent
    bad value is worse than a loud one.
    """
    warnings = []
    for i, row in enumerate(rows):
        label = row.get("Vendor") or f"row {i}"

        ar = str(row.get("Auto-Renew", "")).strip()
        if ar and ar not in AUTO_RENEW_VALUES:
            warnings.append(
                f"{label}: Auto-Renew '{ar}' is not one of {sorted(AUTO_RENEW_VALUES)}"
            )

        st = str(row.get("Status", "")).strip()
        if st and st not in STATUS_VALUES:
            warnings.append(
                f"{label}: Status '{st}' is not one of {sorted(STATUS_VALUES)}"
            )

        cat = str(row.get("Category", "")).strip()
        if cat and cat not in CATEGORIES:
            warnings.append(
                f"{label}: Category '{cat}' is not a standard bucket {sorted(CATEGORIES)}"
            )

        # Next Event semantics: blank when the sub has Ended.
        if st == "Ended" and str(row.get("Next Event", "")).strip():
            warnings.append(
                f"{label}: Status is Ended but Next Event is set — it should be blank"
            )

        # Amount must be a number or the literal VERIFY, never a guess dressed
        # as prose.
        amt = str(row.get("Amount", "")).strip()
        if amt and amt != "VERIFY":
            cleaned = amt.lstrip("$").replace(",", "")
            try:
                float(cleaned)
            except ValueError:
                warnings.append(
                    f"{label}: Amount '{amt}' is neither a number nor 'VERIFY'"
                )
    return warnings


def monthly_equivalent(amount, cycle):
    """Best-effort monthly value for spend totals; returns None if not summable.

    Only Auto-Renew On rows should be passed here by the caller. VERIFY amounts
    and unparseable cycles return None so they are excluded from totals rather
    than guessed.
    """
    amt = str(amount).strip()
    if not amt or amt == "VERIFY":
        return None
    try:
        value = float(amt.lstrip("$").replace(",", ""))
    except ValueError:
        return None
    c = str(cycle).strip().lower()
    if c in ("monthly", "month", "mo", "1 month"):
        return value
    if c in ("yearly", "annual", "annually", "year", "yr", "1 year"):
        return value / 12.0
    if c in ("quarterly", "quarter", "3 months"):
        return value / 3.0
    if c in ("weekly", "week"):
        return value * 52.0 / 12.0
    if c in ("biannual", "semiannual", "semi-annual", "6 months"):
        return value / 6.0
    return None


def spend_totals(rows):
    """Monthly and annualized spend, counting ONLY Auto-Renew On rows."""
    monthly = 0.0
    counted = 0
    skipped = 0
    for row in rows:
        if str(row.get("Auto-Renew", "")).strip() != "On":
            continue
        m = monthly_equivalent(row.get("Amount"), row.get("Billing Cycle"))
        if m is None:
            skipped += 1
            continue
        monthly += m
        counted += 1
    return {
        "monthly": round(monthly, 2),
        "annual": round(monthly * 12.0, 2),
        "counted": counted,
        "skipped_unknown_amount": skipped,
    }


def write_xlsx(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Subscriptions"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])
    ws.freeze_panes = "A2"
    # Reasonable default widths.
    for idx, col in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = max(12, len(col) + 2)
    wb.save(path)


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: row.get(col, "") for col in COLUMNS})


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--rows",
        help="Path to a JSON array of subscription rows. Reads stdin if omitted.",
    )
    parser.add_argument(
        "--date",
        required=True,
        help="Snapshot date YYYY-MM-DD (drives the output filename).",
    )
    parser.add_argument(
        "--outdir", default=".", help="Directory to write the snapshot into."
    )
    args = parser.parse_args(argv)

    raw = open(args.rows, encoding="utf-8").read() if args.rows else sys.stdin.read()
    try:
        rows = json.loads(raw) if raw.strip() else []
    except json.JSONDecodeError as exc:
        parser.error(f"--rows is not valid JSON: {exc}")
    if not isinstance(rows, list):
        parser.error("rows JSON must be a top-level array of objects")

    warnings = validate(rows)
    for w in warnings:
        print(f"WARNING: {w}", file=sys.stderr)

    base = f"Subscription Tracker {args.date}"
    outdir = args.outdir.rstrip("/") or "."

    try:
        path = f"{outdir}/{base}.xlsx"
        write_xlsx(path, rows)
    except ImportError:
        path = f"{outdir}/{base}.csv"
        write_csv(path, rows)
        print(
            "openpyxl not installed — wrote CSV instead of XLSX.", file=sys.stderr
        )

    totals = spend_totals(rows)
    print(f"Wrote {path} ({len(rows)} rows)")
    print(
        f"Go-forward spend (Auto-Renew On only): "
        f"${totals['monthly']}/mo, ${totals['annual']}/yr "
        f"across {totals['counted']} subs "
        f"({totals['skipped_unknown_amount']} excluded for unknown amount)"
    )
    if warnings:
        print(f"{len(warnings)} data warning(s) — see stderr above.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
